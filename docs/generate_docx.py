import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import re
import os

def create_report_docx():
    doc = docx.Document()

    # Set standard margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Color definitions
    NAVY = RGBColor(15, 23, 42)
    BRAND_PURPLE = RGBColor(124, 58, 237)
    MUTED_GRAY = RGBColor(100, 116, 139)

    # Helper styles
    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(24)
        run.font.color.rgb = BRAND_PURPLE
        p.paragraph_format.space_after = Pt(8)

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = NAVY
        p.paragraph_format.space_after = Pt(24)

    def add_heading_1(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(16)
        run.font.color.rgb = BRAND_PURPLE
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)

    def add_heading_2(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(13)
        run.font.color.rgb = NAVY
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)

    def add_para(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.font.size = Pt(10.5)
        return p

    def add_bullet(text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.font.size = Pt(10.5)

    md_path = os.path.join(os.path.dirname(__file__), 'College_Project_Report.md')
    if not os.path.exists(md_path):
        print(f"Error: {md_path} not found")
        return

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    add_title("EVENTHUB – EVENT DISCOVERY & REGISTRATION PLATFORM")
    add_subtitle("A Complete MERN Stack Academic Capstone Project Report")

    in_code_block = False
    in_table = False
    table_rows = []

    for line in lines:
        stripped = line.strip()

        # Handle code blocks
        if stripped.startswith('```'):
            if in_code_block:
                in_code_block = False
            else:
                in_code_block = True
            continue

        if in_code_block:
            # Add as monospace small text
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(line.rstrip())
            run.font.size = Pt(9)
            run.font.name = 'Consolas'
            run.font.color.rgb = RGBColor(51, 65, 85)
            continue

        # Handle Table
        if '|' in stripped and stripped.startswith('|') and stripped.endswith('|'):
            # Table row
            if re.match(r'^\|[\s\-:|]+\|$', stripped):
                # separator line
                continue
            cells = [c.strip() for c in stripped.split('|')[1:-1]]
            table_rows.append(cells)
            in_table = True
            continue
        else:
            if in_table and table_rows:
                # Flush table to document
                cols = len(table_rows[0])
                table = doc.add_table(rows=len(table_rows), cols=cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.style = 'Light Grid Accent 1' if 'Light Grid Accent 1' in [s.name for s in doc.styles] else 'Table Grid'
                for r_idx, row in enumerate(table_rows):
                    for c_idx, cell_text in enumerate(row):
                        if c_idx < cols:
                            cell = table.cell(r_idx, c_idx)
                            cell.text = cell_text
                            # Format header row
                            if r_idx == 0:
                                for p in cell.paragraphs:
                                    for run in p.runs:
                                        run.bold = True
                                        run.font.size = Pt(9.5)
                            else:
                                for p in cell.paragraphs:
                                    for run in p.runs:
                                        run.font.size = Pt(9)
                doc.add_paragraph().paragraph_format.space_after = Pt(6)
                table_rows = []
                in_table = False

        if not stripped:
            continue

        # Headings
        if stripped.startswith('# '):
            continue # already added title
        elif stripped.startswith('## '):
            add_heading_1(stripped[3:])
        elif stripped.startswith('### '):
            add_heading_2(stripped[4:])
        elif stripped.startswith('#### '):
            add_heading_2(stripped[5:])
        elif stripped.startswith('- ') or stripped.startswith('* '):
            add_bullet(stripped[2:])
        elif re.match(r'^\d+\.\s', stripped):
            # Numbered list
            text = re.sub(r'^\d+\.\s', '', stripped)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(text)
            run.font.size = Pt(10.5)
        else:
            # Normal paragraph (clean basic markdown asterisks)
            clean_text = stripped.replace('**', '').replace('__', '').replace('`', '')
            add_para(clean_text)

    # If any remaining table rows
    if table_rows:
        cols = len(table_rows[0])
        table = doc.add_table(rows=len(table_rows), cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r_idx, row in enumerate(table_rows):
            for c_idx, cell_text in enumerate(row):
                if c_idx < cols:
                    cell = table.cell(r_idx, c_idx)
                    cell.text = cell_text

    out_path = os.path.join(os.path.dirname(__file__), 'College_Project_Report.docx')
    doc.save(out_path)
    print(f"Successfully generated academic project report at: {out_path}")

if __name__ == '__main__':
    create_report_docx()
